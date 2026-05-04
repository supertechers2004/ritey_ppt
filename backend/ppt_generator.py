import os
import random
import requests
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from io import BytesIO
import urllib.parse

class PPTGenerator:
    def __init__(self, title='presentation.pptx', theme_name=None):
        """Initialize the presentation generator"""
        self.title = title

        self.FONT_SIZES = {
            'title': 44,
            'intro': 20,
            'bullets': 18,
            'sub_bullets': 16,
            'supporting': 16,
            'paragraphs': 16,
            'small_text': 14
        }

        template_folder = r"templates"

        if os.path.exists(template_folder):
            templates = [f for f in os.listdir(template_folder) if f.endswith(".pptx")]

            if templates:
                if theme_name and theme_name in templates:
                    selected_template = theme_name
                else:
                    selected_template = random.choice(templates)
                self.prs = Presentation(os.path.join(template_folder, selected_template))
            else:
                self.prs = Presentation()
        else:
            self.prs = Presentation()

        # Remove all existing slides from the template properly
        while len(self.prs.slides) > 0:
            rId = self.prs.slides._sldIdLst[0].rId
            self.prs.part.drop_rel(rId)
            del self.prs.slides._sldIdLst[0]

    def _fetch_pexels_image(self, query: str) -> tuple[str | None, bytes | None]:
        """Search Pexels for a landscape image matching the query.
        Returns (image_url, image_bytes) or (None, None) on failure."""
        try:
            pexels_key = os.getenv("PEXELS_API")
            if not pexels_key:
                return None, None

            search_url = (
                f"https://api.pexels.com/v1/search"
                f"?query={urllib.parse.quote(query)}&per_page=1&orientation=landscape"
            )
            res = requests.get(
                search_url,
                headers={"Authorization": pexels_key},
                timeout=10
            )
            if res.status_code != 200:
                return None, None

            photos = res.json().get("photos", [])
            if not photos:
                return None, None

            image_url = photos[0]["src"]["large"]
            img_res = requests.get(image_url, timeout=15)
            if img_res.status_code == 200:
                return image_url, img_res.content

        except Exception as e:
            print(f"Error fetching Pexels image: {e}")

        return None, None

    def add_title_slide(self, title, subtitle="", cover_image_bytes=None):
        """Add a title slide using the template's own picture-placeholder layout when available."""
        from pptx.enum.shapes import PP_PLACEHOLDER

        try:
            # Collect ALL layouts that have a picture placeholder, then pick one randomly
            picture_layouts = []
            if cover_image_bytes:
                for layout in self.prs.slide_layouts:
                    for ph in layout.placeholders:
                        if ph.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
                            picture_layouts.append(layout)
                            break  # one match per layout is enough

            # Random variety on every generation; fall back to plain layout 0 if none found
            layout = random.choice(picture_layouts) if picture_layouts else self.prs.slide_layouts[0]
            slide = self.prs.slides.add_slide(layout)

            # Fill each placeholder by its semantic type
            for ph in slide.placeholders:
                ph_type = ph.placeholder_format.type

                if ph_type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                    ph.text = title or "Presentation"
                    for para in ph.text_frame.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(self.FONT_SIZES["title"])
                            run.font.bold = True

                elif ph_type in (PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
                    ph.text = subtitle

                elif ph_type == PP_PLACEHOLDER.PICTURE and cover_image_bytes:
                    try:
                        ph.insert_picture(BytesIO(cover_image_bytes))
                    except Exception as e:
                        print(f"Error inserting picture placeholder: {e}")

            return slide
        except Exception as e:
            print(f"Error adding title slide: {e}")
            return self.prs.slides.add_slide(self.prs.slide_layouts[-1])


    def add_content_slide(self, slide_data):
        """Add a content slide — text only, no per-slide images."""
        try:
            layouts = self.prs.slide_layouts
            layout = layouts[1] if len(layouts) > 1 else layouts[0]
            slide = self.prs.slides.add_slide(layout)

            # ===== TITLE =====
            title_text = slide_data.get("slide_title", "Untitled Slide")
            if slide.shapes.title:
                slide.shapes.title.text = title_text
                for paragraph in slide.shapes.title.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(self.FONT_SIZES["title"])

            # ===== FIND BODY PLACEHOLDER =====
            body_shape = None
            for shape in slide.placeholders:
                if shape.placeholder_format.type in (2, 7):
                    body_shape = shape
                    break
            if not body_shape:
                for shape in slide.placeholders:
                    if shape != slide.shapes.title:
                        body_shape = shape
                        break
            if not body_shape:
                return slide

            text_frame = body_shape.text_frame
            text_frame.clear()
            text_frame.word_wrap = True

            # ===== INTRO LINE =====
            intro = slide_data.get("intro_line")
            if intro:
                p = text_frame.paragraphs[0]
                p.text = intro
                p.level = 0
                p.space_after = Pt(8)
                for run in p.runs:
                    run.font.size = Pt(self.FONT_SIZES["intro"])

            # ===== BULLET POINTS =====
            bullets = slide_data.get("bullet_points", [])
            for bullet in bullets:
                p = text_frame.add_paragraph()
                p.text = str(bullet)
                p.level = 1
                p.space_after = Pt(4)
                for run in p.runs:
                    run.font.size = Pt(self.FONT_SIZES["bullets"])

            # ===== SUPPORTING TEXT =====
            supporting = slide_data.get("supporting_text")
            if supporting:
                text_frame.add_paragraph()  # spacer
                p = text_frame.add_paragraph()
                p.text = str(supporting)
                p.level = 0
                p.space_after = Pt(6)
                for run in p.runs:
                    run.font.size = Pt(self.FONT_SIZES["supporting"])
                    run.font.italic = True

            # ===== PARAGRAPHS =====
            paragraphs = slide_data.get("paragraphs", [])
            for para in paragraphs:
                p = text_frame.add_paragraph()
                p.text = str(para)
                p.level = 0
                p.space_after = Pt(8)
                for run in p.runs:
                    run.font.size = Pt(self.FONT_SIZES["paragraphs"])

            return slide
        except Exception as e:
            print(f"Error adding content slide: {e}")
            return None

    def generate_from_list(self, slides_data) -> str | None:
        """Build the full presentation. Returns the Pexels cover image URL (or None)."""

        # Determine a good cover image query from the first slide or the title
        cover_query = self.title
        if slides_data:
            cover_query = slides_data[0].get("image_query") or self.title

        cover_image_url, cover_image_bytes = self._fetch_pexels_image(cover_query)

        # Title slide — with cover image
        self.add_title_slide(self.title, "Generated by Ritey AI", cover_image_bytes=cover_image_bytes)

        # Content slides — text only
        for slide in slides_data:
            self.add_content_slide(slide)

        # Closing slide — no image
        self.add_title_slide("Thank You!", "Questions?")

        return cover_image_url  # caller stores this in DB

    def save(self):
        ppt_io = BytesIO()
        self.prs.save(ppt_io)
        ppt_io.seek(0)
        return ppt_io
